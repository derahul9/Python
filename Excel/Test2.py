from openpyxl import Workbook
from openpyxl.styles import *
from openpyxl.worksheet.table import Table, TableStyleInfo

text_file = open("employees.txt")

records = []
text_file.seek(0)

for record in text_file.readlines():
    records.append(record.rstrip("\n").split(";"))

print (records)

workbook = Workbook()
file = "MyCompanyStaff.xlsx"
workbook.save(file)

print(workbook.sheetnames)
sheet = workbook['Sheet']
sheet.title = 'Employees'

for row in records:
    sheet.append(row)

table = table(displayname = Table, ref = "A1:G11")

# Defining a style for the table (default style name, row/column stripes)
# Choose your table style from the default styles of openpyxl
# Just type in openpyxl.worksheet.table.TABLESTYLES in the Python interpreter
style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=True)

# Applying the style to the table
table.tableStyleInfo = style

# Adding the newly created table to the sheet
sheet.add_table(table)

# Defining the font (red, bold, italic) for salary > 55000
font = Font(color=colors.RED, bold=True, italic=True)

# Applying the font settings to the cells that meet the condition salary > 55000
for cell_no in range(2, 12):
    if int(sheet['G%s' % (cell_no)].value) > 55000:
        sheet['G%s' % (cell_no)].font = font

# Saving the changes made to the workbook
workbook.save(file_path)

# Closing the text file
text_file.close()

# Closing the workbook
